#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script to generate multi-sheet Excel export based on Source/VMC Form 01.xlsx
Sheet format: 'แบบฟอร์มยานพาหนะใหม่'
Each selected vehicle gets its own sheet. All sheets in 1 Excel workbook.
"""

import sys
import io
import os
import re
import json
import argparse
import calendar
from datetime import datetime
import openpyxl

# Force UTF-8 stdout
if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

THAI_MONTHS = [
    '', 'มกราคม', 'กุมภาพันธ์', 'มีนาคม', 'เมษายน', 'พฤษภาคม', 'มิถุนายน',
    'กรกฎาคม', 'สิงหาคม', 'กันยายน', 'ตุลาคม', 'พฤศจิกายน', 'ธันวาคม'
]

def sanitize_sheet_name(name, existing_names):
    # Excel sheet title restrictions: max 31 chars, no : \ / ? * [ ]
    clean = re.sub(r'[\/?:*\[\]\\]', '-', str(name or 'รถ')).strip()
    if not clean:
        clean = 'ยานพาหนะ'
    clean = clean[:28]
    candidate = clean
    counter = 1
    while candidate in existing_names:
        candidate = f"{clean[:25]}_{counter}"
        counter += 1
    return candidate

def get_days_in_month(year_str, month_str):
    try:
        y = int(year_str)
        m = int(month_str)
        return calendar.monthrange(y, m)[1]
    except Exception:
        return 30

def export_vmc_excel(db_path, template_path, output_path, month_str, vehicle_ids=None):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template not found at: {template_path}")
    
    with open(db_path, 'r', encoding='utf-8') as f:
        db = json.load(f)

    vehicles_list = db.get('vehicles', [])
    pn1_reports = db.get('pn1_reports', [])

    # Filter vehicles
    if vehicle_ids and vehicle_ids != ['all'] and 'all' not in vehicle_ids:
        v_ids_set = {int(x) for x in vehicle_ids if str(x).isdigit()}
        selected_vehicles = [v for v in vehicles_list if v.get('id') in v_ids_set]
    else:
        selected_vehicles = vehicles_list

    if not selected_vehicles:
        raise ValueError("No vehicles found to export")

    # Parse Month and Year (BE)
    try:
        parts = month_str.split('-')
        year_ce = int(parts[0])
        month_num = int(parts[1])
    except Exception:
        year_ce = 2026
        month_num = 9

    year_be = year_ce + 543
    month_name_th = THAI_MONTHS[month_num] if 1 <= month_num <= 12 else str(month_num)
    days_in_month = get_days_in_month(year_ce, month_num)

    # Load template workbook
    wb = openpyxl.load_workbook(template_path)
    if 'แบบฟอร์มยานพาหนะใหม่' not in wb.sheetnames:
        raise ValueError("Sheet 'แบบฟอร์มยานพาหนะใหม่' not found in template")
    
    tmpl_ws = wb['แบบฟอร์มยานพาหนะใหม่']

    used_sheet_names = set()

    for vehicle in selected_vehicles:
        v_id = vehicle.get('id')
        internal_id = vehicle.get('internal_id') or f"รถลำดับ-{v_id}"
        license_plate = vehicle.get('license_plate') or ''
        brand_model = f"{vehicle.get('brand', '')} {vehicle.get('model', '')}".strip() or vehicle.get('category', '')
        department = vehicle.get('department') or 'อหก.'
        fuel_type = vehicle.get('fuel_type') or 'ดีเซล'

        # Find PN1 report for this vehicle & month
        report = next(
            (r for r in pn1_reports if r.get('vehicle_id') == v_id and r.get('month') == month_str),
            None
        )

        daily_logs_map = {}
        if report and report.get('daily_logs'):
            for log in report.get('daily_logs', []):
                d_str = log.get('date', '')
                if d_str:
                    try:
                        day_num = int(d_str.split('-')[2])
                        daily_logs_map[day_num] = log
                    except Exception:
                        pass

        sheet_title = sanitize_sheet_name(internal_id, used_sheet_names)
        used_sheet_names.add(sheet_title)

        # Clone template worksheet
        ws = wb.copy_worksheet(tmpl_ws)
        ws.title = sheet_title

        # Header population
        # A3: Department header
        dept_text = f"ฝ่ายบริหารและก่อสร้างโรงไฟฟ้า (อหก.) {department}"
        ws['A3'] = dept_text

        # Row 4: ทะเบียน กฟผ. (E4), ชนิดน้ำมันเชื้อเพลิง (Q4)
        ws['E4'] = internal_id
        ws['Q4'] = fuel_type

        # Row 5: ทะเบียนขนส่ง (E5), หน่วยงาน (Q5)
        ws['E5'] = license_plate
        ws['Q5'] = department

        # Row 6: ยี่ห้อและรุ่นของรถ (E6), ประจำเดือน (P6), พ.ศ. (S6)
        ws['E6'] = brand_model
        ws['P6'] = month_name_th
        ws['S6'] = str(year_be)

        # Populate rows 12 to 42 (Days 1 to 31)
        for day in range(1, 32):
            row_idx = 11 + day  # day 1 -> row 12
            ws.cell(row_idx, 1).value = day  # Col A: วันที่

            if day > days_in_month:
                # Clear content for non-existent days in this month
                ws.cell(row_idx, 2).value = None
                ws.cell(row_idx, 3).value = None
                ws.cell(row_idx, 4).value = None
                ws.cell(row_idx, 5).value = None
                for c in range(6, 14):
                    ws.cell(row_idx, c).value = None
                ws.cell(row_idx, 14).value = None
                ws.cell(row_idx, 15).value = None
                ws.cell(row_idx, 17).value = None
                ws.cell(row_idx, 18).value = None
                ws.cell(row_idx, 19).value = None
                continue

            log = daily_logs_map.get(day)

            if log:
                code = str(log.get('status_code', '/')).strip()
                if code == '/' or code == 'ใช้งาน':
                    status_text = 'ใช้งาน'
                elif code == '0' or code == 'ซ่อม':
                    status_text = 'ซ่อม'
                else:
                    status_text = 'ว่างงาน'

                ws.cell(row_idx, 2).value = status_text

                # Mileages
                if log.get('start_mileage') is not None and log.get('start_mileage') != '':
                    ws.cell(row_idx, 3).value = float(log.get('start_mileage'))
                if log.get('end_mileage') is not None and log.get('end_mileage') != '':
                    ws.cell(row_idx, 4).value = float(log.get('end_mileage'))
                if log.get('fuel_mileage') is not None and log.get('fuel_mileage') != '':
                    ws.cell(row_idx, 5).value = float(log.get('fuel_mileage'))

                # Checklist (Cols F to M: cols 6 to 13)
                chk = log.get('checklist') or {}
                checklist_keys = [
                    'coolant', 'engine_oil', 'brake', 'clutch',
                    'tire', 'signal_light', 'hydraulic', 'vehicle_body'
                ]
                for c_offset, key in enumerate(checklist_keys):
                    col_num = 6 + c_offset
                    val = chk.get(key)
                    # If true, mark ✓
                    if val is True:
                        ws.cell(row_idx, col_num).value = '✓'
                    elif val is False:
                        ws.cell(row_idx, col_num).value = ''

                # Refuel info
                liters = log.get('fuel_liters')
                price = log.get('fuel_price_per_liter')
                if liters and float(liters) > 0:
                    ws.cell(row_idx, 14).value = float(liters)
                    if price and float(price) > 0:
                        ws.cell(row_idx, 15).value = float(price)

                # Department, Driver, Destination
                ws.cell(row_idx, 17).value = log.get('department') or department
                ws.cell(row_idx, 18).value = log.get('driver_name') or ''
                ws.cell(row_idx, 19).value = log.get('destination') or ''

            else:
                # Default for days without trip: 'ว่างงาน'
                ws.cell(row_idx, 2).value = 'ว่างงาน'

        # Footer row 44: Working days in this month
        ws['E44'] = days_in_month

        # Configure page setup for 1-page A4 print
        ws.print_area = 'A1:S49'
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4

    # Remove the 2 base template sheets so only actual vehicle sheets remain
    if 'แบบฟอร์มรายงานยานพาหนะ' in wb.sheetnames:
        wb.remove(wb['แบบฟอร์มรายงานยานพาหนะ'])
    if 'แบบฟอร์มยานพาหนะใหม่' in wb.sheetnames:
        wb.remove(wb['แบบฟอร์มยานพาหนะใหม่'])

    # Save workbook
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return {
        'success': True,
        'vehicles_count': len(selected_vehicles),
        'sheet_names': wb.sheetnames,
        'output_path': output_path
    }

def main():
    parser = argparse.ArgumentParser(description='Export PN1 reports to VMC Form 01 Excel')
    parser.add_argument('--db', default='server/database.json', help='Path to database.json')
    parser.add_argument('--template', default='Source/VMC Form 01.xlsx', help='Path to VMC Form 01.xlsx')
    parser.add_argument('--month', default='2026-09', help='Report month YYYY-MM')
    parser.add_argument('--vehicles', default='all', help='Comma separated vehicle IDs or "all"')
    parser.add_argument('--output', default='Source/temp_export.xlsx', help='Output file path')

    args = parser.parse_args()

    v_ids = args.vehicles.split(',') if args.vehicles != 'all' else ['all']
    result = export_vmc_excel(
        db_path=args.db,
        template_path=args.template,
        output_path=args.output,
        month_str=args.month,
        vehicle_ids=v_ids
    )

    print(json.dumps(result, ensure_ascii=False))

if __name__ == '__main__':
    main()
