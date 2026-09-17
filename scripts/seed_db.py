import os
import sys
import io
import datetime
import json
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

JSON_PATH = os.path.join(os.path.dirname(__file__), '..', 'server', 'database.json')
EXCEL_PATH = os.path.join(os.path.dirname(__file__), '..', 'Source', 'Database.xlsx')

os.makedirs(os.path.dirname(JSON_PATH), exist_ok=True)

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
sheet = wb['Database']

def parse_date(val):
    if not val or str(val).strip().upper() == 'NA':
        return None
    if isinstance(val, (datetime.date, datetime.datetime)):
        year = val.year
        if year > 2400: # BE Year
            year -= 543
        return f'{year:04d}-{val.month:02d}-{val.day:02d}'
    if isinstance(val, str):
        parts = val.strip().split('/')
        if len(parts) == 3:
            day, month, year = int(parts[0]), int(parts[1]), int(parts[2])
            if year > 2400:
                year -= 543
            return f'{year:04d}-{month:02d}-{day:02d}'
    return str(val)

def determine_fuel_type(brand, model, vehicle_type):
    b = str(brand or '').upper()
    m = str(model or '').upper()
    vt = str(vehicle_type or '')

    if 'จักรยานยนต์' in vt or 'WAVE' in m or 'HONDA' in b:
        return 'แก๊สโซฮอล์ 95'
    elif 'SEMI-TRAILER' in m or 'VOLVO' in b or 'HINO' in b or 'ISUZU' in b:
        return 'ดีเซล B7'
    elif 'TOYOTA' in b or 'COMMUTER' in m:
        return 'ดีเซล'
    elif 'BENZ' in b or 'MAZDA' in b or 'NISSAN' in b:
        return 'ดีเซล B7'
    return 'ดีเซล'

# Balanced status distribution for testing features:
# 20 Standby ("จอดรองาน"), 25 In-Use ("ใช้งาน"), 6 Maintenance ("รอซ่อม")
statuses = ['จอดรองาน'] * 20 + ['ใช้งาน'] * 25 + ['รอซ่อม'] * 6

inserted_vehicles = []

for r in range(2, sheet.max_row + 1):
    seq = sheet.cell(r, 1).value
    if seq is None:
        continue
    overall_dept = str(sheet.cell(r, 2).value or 'อหก.').strip()
    internal_id = str(sheet.cell(r, 3).value or '').strip()
    license_plate = str(sheet.cell(r, 4).value or '').strip()
    brand = str(sheet.cell(r, 5).value or '').strip()
    model = str(sheet.cell(r, 6).value or '').strip()
    category = str(sheet.cell(r, 7).value or '').strip()
    vehicle_type = str(sheet.cell(r, 8).value or '').strip()
    specs = str(sheet.cell(r, 9).value or '').strip() if sheet.cell(r, 9).value else ''
    
    # Col 10: ชนิดของเชื้อเพลิง
    raw_fuel = str(sheet.cell(r, 10).value or '').strip()
    if not raw_fuel or raw_fuel.upper() == 'NA':
        fuel_type = determine_fuel_type(brand, model, vehicle_type)
    else:
        fuel_type = raw_fuel
    
    # Col 11 (K): เลขไมล์ล่าสุด
    raw_mileage = sheet.cell(r, 11).value
    try:
        mileage = float(raw_mileage) if raw_mileage is not None and str(raw_mileage).strip() != '' else 0.0
    except (ValueError, TypeError):
        mileage = 0.0

    # Col 12 (L): Base Location
    raw_location = str(sheet.cell(r, 12).value or '').strip()
    base_location = raw_location if raw_location else 'สนก.บางกรวย'
    
    # Col 13 (M): สังกัด (หน่วยงานย่อย/เจ้าของรถ เช่น กยค-พ., กคร-พ., กฟค-พ.)
    department = str(sheet.cell(r, 13).value or '').strip()
    if not department or department.upper() == 'NA':
        department = overall_dept

    # Col 14 (N): สถานะ
    excel_status = str(sheet.cell(r, 14).value or '').strip()
    status = excel_status if excel_status in ['จอดรองาน', 'ใช้งาน', 'รอซ่อม', 'ซ่อม', 'รอยุบสภาพ', 'ยุบสภาพ'] else statuses[len(inserted_vehicles) % len(statuses)]
    
    # Col 15 (O): วันจดทะเบียน
    reg_date = parse_date(sheet.cell(r, 15).value)
    
    # Col 16 (P): รอบชำระภาษีประจำปี
    
    # Col 17 (Q): วันครบกำหนดชำระภาษี
    tax_due_date = parse_date(sheet.cell(r, 17).value)
    
    # Col 18 (R): ยอดชำระภาษี 2569
    tax_amt_val = sheet.cell(r, 18).value
    tax_amount = float(tax_amt_val) if tax_amt_val is not None and str(tax_amt_val).strip() != '' and str(tax_amt_val).upper() != 'NONE' else None

    v_id = len(inserted_vehicles) + 1
    inserted_vehicles.append({
        'id': v_id,
        'seq': seq,
        'department': department,
        'internal_id': internal_id,
        'license_plate': license_plate,
        'brand': brand,
        'model': model,
        'category': category,
        'vehicle_type': vehicle_type,
        'fuel_type': fuel_type,
        'specifications': specs,
        'base_location': base_location,
        'status': status,
        'registration_date': reg_date,
        'tax_due_date': tax_due_date,
        'tax_amount': tax_amount,
        'current_mileage': mileage,
        'created_at': datetime.datetime.now().isoformat()
    })

work_orders = []
in_use_vehicles = [v for v in inserted_vehicles if v['status'] == 'ใช้งาน']
for idx, v in enumerate(in_use_vehicles[:10]):
    work_orders.append({
        'id': len(work_orders) + 1,
        'vehicle_id': v['id'],
        'requester_name': f'สมชาย ใจดี {idx+1}',
        'department': 'แผนกบำรุงรักษา อหก.',
        'purpose': 'ออกตรวจเช็คสถานีไฟฟ้าและระบบเครื่องจักรประจำสายงาน',
        'start_datetime': '2026-09-14 08:30:00',
        'end_datetime': None,
        'start_mileage': v['current_mileage'] - 120,
        'end_mileage': None,
        'total_distance': None,
        'status': 'IN_PROGRESS',
        'created_at': '2026-09-14 08:30:00'
    })

for idx, v in enumerate(inserted_vehicles[:8]):
    start_m = v['current_mileage'] - 450
    end_m = v['current_mileage'] - 120
    work_orders.append({
        'id': len(work_orders) + 1,
        'vehicle_id': v['id'],
        'requester_name': f'วิชัย การช่าง {idx+1}',
        'department': 'แผนกปฏิบัติการ อหก.',
        'purpose': 'เดินทางไปปฏิบัติงาน สนก.บางกรวย',
        'start_datetime': '2026-09-10 09:00:00',
        'end_datetime': '2026-09-10 17:00:00',
        'start_mileage': start_m,
        'end_mileage': end_m,
        'total_distance': end_m - start_m,
        'status': 'COMPLETED',
        'created_at': '2026-09-10 09:00:00'
    })

maintenance_records = []
for idx, v in enumerate(inserted_vehicles[::5]):
    maintenance_records.append({
        'id': len(maintenance_records) + 1,
        'vehicle_id': v['id'],
        'record_type': 'ตรอ. / ถ่ายน้ำมันเครื่อง',
        'service_date': '2026-06-15',
        'next_due_date': '2026-12-15',
        'cost': 3500.0,
        'remarks': 'เปลี่ยนถ่ายน้ำมันเครื่อง ไส้กรอง และตรวจเช็คระบบเบรกประจำปี',
        'created_at': '2026-06-15 10:00:00'
    })

fuel_logs = []
for idx, v in enumerate(inserted_vehicles[::3]):
    fuel_logs.append({
        'id': len(fuel_logs) + 1,
        'vehicle_id': v['id'],
        'work_order_id': None,
        'refuel_date': f'2026-09-{(idx%10)+1:02d} 10:15:00',
        'liters': 45.5,
        'total_amount': 1550.0,
        'odometer': v['current_mileage'] - 50,
        'receipt_no': f'INV-2026-{1000+idx}',
        'created_at': f'2026-09-{(idx%10)+1:02d} 10:15:00'
    })

db_data = {
    'vehicles': inserted_vehicles,
    'work_orders': work_orders,
    'maintenance_records': maintenance_records,
    'fuel_logs': fuel_logs,
    'email_logs': []
}

with open(JSON_PATH, 'w', encoding='utf-8') as f:
    json.dump(db_data, f, ensure_ascii=False, indent=2)

print(f"Successfully created JSON database at {JSON_PATH}")
print(f"Inserted {len(inserted_vehicles)} vehicles with department and fuel_type.")
